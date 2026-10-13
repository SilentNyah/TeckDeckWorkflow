# -*- coding: utf-8 -*-
"""Tests for the step-F additions to pptx-toolkit.

    python -m unittest discover -s scripts/tests -t scripts      (from skills/pptx-toolkit)

Covers: textbox(name=), slide_fingerprints / assert_untouched / save_replace,
check_deck's visual.uncounted and the R8-7 wording, phrase_breaks / wrap_ja(phrase=) and
the R8-2 phrase-break lint (the last three skip when budoux is not installed).
"""
import os
import sys
import tempfile
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.dirname(HERE))

import deckkit  # noqa: E402
import check_deck  # noqa: E402

try:
    import budoux  # noqa: F401
    HAVE_BUDOUX = True
except ImportError:
    HAVE_BUDOUX = False

CJK = 'Yu Gothic'


def make_deck(path, n=3):
    prs = deckkit.new_presentation()
    for _ in range(n):
        deckkit.add_slide(prs, bg='151124')
    prs.save(path)
    return prs


class TestNamesAndGuards(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.mkdtemp(prefix='toolkit-')
        self.deck = os.path.join(self.dir, 'deck.pptx')
        make_deck(self.deck)

    def test_textbox_name_prefix_and_default(self):
        from pptx import Presentation
        prs = Presentation(self.deck)
        s = prs.slides[0]
        el = deckkit.textbox(s, 1, 1, 3, 0.5, [[deckkit.R('出典', 7, 'ECEAF7', CJK)]], name='src:footer')
        el2 = deckkit.textbox(s, 1, 2, 3, 0.5, [[deckkit.R('本文', 12, 'ECEAF7', CJK)]])
        names = [sh.name for sh in s.shapes]
        self.assertTrue(any(nm.startswith('src:footer') for nm in names), names)
        self.assertTrue(any(nm.startswith('TextBox') for nm in names), names)
        self.assertIsNotNone(el)
        self.assertIsNotNone(el2)

    def test_fingerprints_detect_change_outside_target(self):
        from pptx import Presentation
        before = deckkit.slide_fingerprints(self.deck)
        self.assertEqual(sorted(before), [1, 2, 3])
        prs = Presentation(self.deck)
        deckkit.textbox(prs.slides[2], 1, 1, 3, 0.5, [[deckkit.R('x', 12, 'ECEAF7', CJK)]])
        deckkit.save_replace(prs, self.deck)
        self.assertFalse(os.path.exists(self.deck + '.tmp-%d' % os.getpid()))
        after = deckkit.slide_fingerprints(self.deck)
        with self.assertRaises(AssertionError) as cm:
            deckkit.assert_untouched(before, after, touched=[2])
        self.assertIn('[3]', str(cm.exception))
        self.assertTrue(deckkit.assert_untouched(before, after, touched=[2, 3]))

    def test_fingerprints_detect_count_change(self):
        from pptx import Presentation
        before = deckkit.slide_fingerprints(self.deck)
        prs = Presentation(self.deck)
        deckkit.add_slide(prs, bg='151124')
        prs.save(self.deck)
        with self.assertRaises(AssertionError):
            deckkit.assert_untouched(before, deckkit.slide_fingerprints(self.deck), touched=[1, 2, 3])


class TestCheckDeckBudget(unittest.TestCase):
    def setUp(self):
        self.dir = tempfile.mkdtemp(prefix='toolkit-')
        self.deck = os.path.join(self.dir, 'deck.pptx')
        self.prs = make_deck(self.deck, n=1)
        self.slide = self.prs.slides[0]

    def findings(self, profile):
        out = []
        check_deck.visual_checks(self.prs, 1, self.slide, profile, out)
        return out

    def test_uncounted_prefix_is_excluded_and_reported(self):
        body = '読者が文として読む本文がここに入る。' * 2          # 36 chars
        deckkit.textbox(self.slide, 0.5, 1.0, 8, 1.0, [[deckkit.R(body, 12, 'ECEAF7', CJK)]])
        deckkit.textbox(self.slide, 0.5, 4.5, 8, 0.3,
                        [[deckkit.R('SRC-001 ' * 20, 7, '6E6A8A', CJK)]], name='src:footer')
        deckkit.textbox(self.slide, 0.5, 0.2, 8, 0.4,
                        [[deckkit.R('タイトル' * 10, 18, 'ECEAF7', CJK, b=True)]], name='title:main')
        prof = {'visual': {'body_char_budget': 10}}
        msgs = [m for _, sev, m in self.findings(prof) if '[R8-7]' in m]
        self.assertEqual(len(msgs), 1, msgs)
        self.assertIn('%d characters of body text' % len(body), msgs[0])
        self.assertIn('in uncounted shapes', msgs[0])
        self.assertNotIn('convert some to a diagram', msgs[0])
        self.assertIn('over by %d' % (len(body) - 10), msgs[0])

    def test_unnamed_shapes_all_count(self):
        deckkit.textbox(self.slide, 0.5, 1.0, 8, 1.0, [[deckkit.R('あ' * 30, 12, 'ECEAF7', CJK)]])
        deckkit.textbox(self.slide, 0.5, 4.5, 8, 0.3, [[deckkit.R('い' * 30, 7, '6E6A8A', CJK)]])
        prof = {'visual': {'body_char_budget': 10}}
        msgs = [m for _, _, m in self.findings(prof) if '[R8-7]' in m]
        self.assertIn('60 characters of body text,', msgs[0])

    def test_profile_can_override_uncounted(self):
        deckkit.textbox(self.slide, 0.5, 1.0, 8, 1.0, [[deckkit.R('あ' * 30, 12, 'ECEAF7', CJK)]],
                        name='memo:x')
        prof = {'visual': {'body_char_budget': 10, 'uncounted': ['memo:']}}
        msgs = [m for _, _, m in self.findings(prof) if '[R8-7]' in m]
        self.assertEqual(msgs, [])


@unittest.skipUnless(HAVE_BUDOUX, 'budoux not installed')
class TestPhrases(unittest.TestCase):
    def test_phrase_breaks_keep_brackets_whole(self):
        ph = deckkit.phrase_breaks('システムエンジニアリングサービス（SES＝準委任契約の呼び名）')
        self.assertEqual(''.join(ph), 'システムエンジニアリングサービス（SES＝準委任契約の呼び名）')
        self.assertTrue(any(p.startswith('（') and p.endswith('）') for p in ph), ph)

    def test_wrap_ja_phrase_mode_breaks_on_boundaries(self):
        text = '派遣先はあなたを雇っていないので、派遣ではなく労働者供給になる（→ p.7）'
        # 2.6in holds the longest phrase (9 full-width chars at 18pt = 2.25in)
        lines = deckkit.wrap_ja(text, 2.6, 18, CJK, phrase=True)
        self.assertEqual(''.join(lines), text)
        self.assertGreater(len(lines), 2)
        # a phrase wider than the box falls back to character breaking, without losing text
        tight = deckkit.wrap_ja(text, 2.2, 18, CJK, phrase=True)
        self.assertEqual(''.join(tight), text)
        self.assertTrue(all(tight))
        phrases = deckkit.phrase_breaks(text)
        cuts, pos = set(), 0
        for p in phrases:
            pos += len(p)
            cuts.add(pos)
        pos = 0
        for ln in lines[:-1]:
            pos += len(ln)
            self.assertIn(pos, cuts, 'line break at %d is inside a phrase: %r' % (pos, lines))
        self.assertTrue(any('（→ p.7）' in ln for ln in lines), lines)

    def test_wrap_ja_char_mode_still_works(self):
        text = '派遣先はあなたを雇っていないので、派遣ではなく労働者供給になる'
        lines = deckkit.wrap_ja(text, 2.2, 18, CJK, phrase=False)
        self.assertEqual(''.join(lines), text)
        self.assertGreater(len(lines), 1)

    def test_check_deck_flags_break_inside_phrase(self):
        d = tempfile.mkdtemp(prefix='toolkit-')
        prs = make_deck(os.path.join(d, 'deck.pptx'), n=1)
        s = prs.slides[0]
        deckkit.textbox(s, 0.5, 0.5, 8, 1.0, [[deckkit.R('指示して', 18, 'ECEAF7', CJK, b=True)],
                                              [deckkit.R('いるのは派遣先である', 18, 'ECEAF7', CJK, b=True)]])
        out = []
        check_deck.visual_checks(prs, 1, s, {}, out)
        self.assertTrue(any('[R8-2]' in m for _, _, m in out), out)

    def test_phrase_break_defect_shapes(self):
        pb = deckkit.phrase_breaks
        f = check_deck.phrase_break_defect
        # the three real defects
        self.assertEqual(f('指示して', 'いるのは派遣先である', pb), 'hiragana continuation')
        self.assertEqual(f('取り分にしてしま', 'う会社', pb), 'hiragana continuation')
        self.assertEqual(f('システムエン', 'ジニアリング', pb), 'katakana word split')
        self.assertTrue((f('（SES＝', '準委任契約の呼び名）', pb) or '').startswith('parenthetical'))
        # unrelated lines in one box must NOT fire (seen on real decks)
        self.assertIsNone(f('準委任契約だから', '労働者派遣法は', pb))
        self.assertIsNone(f('画面回転', '構成変更の一例', pb))
        self.assertIsNone(f('「客先に', '常駐していないか」', pb))           # long quote wraps
        self.assertIsNone(f('罰則は行為者個人にも及ぶ。', '次の文。', pb))
        self.assertIsNone(f('罰則は', 'および行政処分', pb))                 # boundary BudouX accepts
        self.assertIsNone(f('プロセス死', 'システム起因', pb))
        self.assertIsNone(f('evaluate(', '"${reservations.size()}")', pb))  # Latin-only / code

    def test_code_paragraphs_are_exempt(self):
        d = tempfile.mkdtemp(prefix='toolkit-')
        prs = make_deck(os.path.join(d, 'deck.pptx'), n=1)
        s = prs.slides[0]
        deckkit.textbox(s, 0.5, 0.5, 8, 1.0,
                        [[deckkit.R('予約一覧（', 9, 'ECEAF7', 'Consolas')],
                         [deckkit.R('件数）を返す', 9, 'ECEAF7', 'Consolas')]])
        out = []
        check_deck.visual_checks(prs, 1, s, {}, out)
        self.assertFalse(any('[R8-2]' in m for _, _, m in out), out)

    def test_check_deck_accepts_break_on_boundary(self):
        d = tempfile.mkdtemp(prefix='toolkit-')
        prs = make_deck(os.path.join(d, 'deck.pptx'), n=1)
        s = prs.slides[0]
        deckkit.textbox(s, 0.5, 0.5, 8, 1.0, [[deckkit.R('指示しているのは', 18, 'ECEAF7', CJK, b=True)],
                                              [deckkit.R('派遣先であり、', 18, 'ECEAF7', CJK, b=True)]])
        deckkit.textbox(s, 0.5, 2.5, 8, 1.0, [[deckkit.R('罰則は行為者個人にも及ぶ。', 12, 'ECEAF7', CJK)],
                                              [deckkit.R('次の文。', 12, 'ECEAF7', CJK)]])
        out = []
        check_deck.visual_checks(prs, 1, s, {}, out)
        self.assertFalse(any('[R8-2]' in m for _, _, m in out), out)


if __name__ == '__main__':
    unittest.main()
