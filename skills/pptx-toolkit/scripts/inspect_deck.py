# -*- coding: utf-8 -*-
"""List the shapes of a deck so you can identify what to edit.

You cannot edit a shape you cannot identify -- this is the tool you run first.

    python inspect_deck.py <deck.pptx>                  # slide count + titles
    python inspect_deck.py <deck.pptx> --slide 39       # every shape on slide 39
    python inspect_deck.py <deck.pptx> --slide 39 --runs   # + per-run formatting
    python inspect_deck.py <deck.pptx> --slide 1-5 --json
"""
import argparse
import json
import sys


def parse_slides(spec, total):
    if not spec:
        return []
    out = []
    for part in str(spec).split(','):
        part = part.strip()
        if '-' in part:
            lo, hi = part.split('-', 1)
            out.extend(range(int(lo), int(hi) + 1))
        else:
            out.append(int(part))
    return [n for n in out if 1 <= n <= total]


def shape_rows(slide):
    from pptx.util import Emu
    rows = []
    for sh in slide.shapes:
        text = ''
        if sh.has_text_frame:
            text = sh.text_frame.text.replace('\n', ' | ')
        try:
            stype = str(sh.shape_type).split()[0]
        except Exception:
            stype = '?'
        rows.append({
            'id': sh.shape_id,
            'type': stype,
            'name': sh.name,
            'x': round(Emu(sh.left).inches, 2) if sh.left is not None else None,
            'y': round(Emu(sh.top).inches, 2) if sh.top is not None else None,
            'w': round(Emu(sh.width).inches, 2) if sh.width is not None else None,
            'h': round(Emu(sh.height).inches, 2) if sh.height is not None else None,
            'text': text,
        })
    return rows


def run_rows(slide):
    out = []
    for sh in slide.shapes:
        if not sh.has_text_frame:
            continue
        for pi, para in enumerate(sh.text_frame.paragraphs):
            for ri, r in enumerate(para.runs):
                f = r.font
                col = None
                try:
                    col = str(f.color.rgb)
                except Exception:
                    pass
                out.append({
                    'shape_id': sh.shape_id, 'p': pi, 'r': ri,
                    'sz': f.size.pt if f.size else None, 'font': f.name,
                    'b': f.bold, 'i': f.italic, 'color': col, 'text': r.text,
                })
    return out


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('deck')
    ap.add_argument('--slide', help='1-based; "39", "1-5", "19,22,25"')
    ap.add_argument('--runs', action='store_true', help='also dump per-run formatting')
    ap.add_argument('--json', action='store_true')
    ap.add_argument('--width', type=int, default=95, help='text truncation width')
    a = ap.parse_args()

    sys.stdout.reconfigure(encoding='utf-8')
    from pptx import Presentation
    from pptx.util import Emu
    prs = Presentation(a.deck)
    total = len(prs.slides)
    targets = parse_slides(a.slide, total)

    if not targets:
        info = {
            'deck': a.deck, 'slides': total,
            'size_in': [round(Emu(prs.slide_width).inches, 3),
                        round(Emu(prs.slide_height).inches, 3)],
            'titles': [],
        }
        for i, s in enumerate(prs.slides, 1):
            t = ''
            for sh in s.shapes:
                if sh.has_text_frame and sh.text_frame.text.strip():
                    t = sh.text_frame.text.strip().splitlines()[0]
                    break
            info['titles'].append({'slide': i, 'title': t})
        if a.json:
            print(json.dumps(info, ensure_ascii=False, indent=2))
        else:
            print('%s\n%d slides, %.3f x %.3f in\n'
                  % (a.deck, total, info['size_in'][0], info['size_in'][1]))
            for row in info['titles']:
                print('  %3d  %s' % (row['slide'], row['title'][:a.width]))
        return

    payload = {}
    for n in targets:
        slide = prs.slides[n - 1]
        payload[n] = {'shapes': shape_rows(slide)}
        if a.runs:
            payload[n]['runs'] = run_rows(slide)

    if a.json:
        print(json.dumps(payload, ensure_ascii=False, indent=2))
        return

    for n in targets:
        print('=' * 24, 'SLIDE', n, '=' * 24)
        for r in payload[n]['shapes']:
            print('  id=%-4s %-11s x=%-6s y=%-6s w=%-6s h=%-6s %s'
                  % (r['id'], r['type'], r['x'], r['y'], r['w'], r['h'],
                     r['text'][:a.width]))
        if a.runs:
            print('  --- runs ---')
            for r in payload[n]['runs']:
                print('    id=%-4s P%sR%s sz=%-5s %-14s b=%-5s %-7s :: %r'
                      % (r['shape_id'], r['p'], r['r'], r['sz'], r['font'],
                         r['b'], r['color'], r['text'][:60]))


if __name__ == '__main__':
    main()
