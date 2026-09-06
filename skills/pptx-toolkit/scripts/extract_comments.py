# -*- coding: utf-8 -*-
"""Extract PowerPoint modern comments and emit a comment-import-rNNN.md skeleton.

Reads raw OOXML from inside the .pptx:
  ppt/comments/modernComment_*.xml    the comments themselves
  ppt/slides/_rels/slideN.xml.rels    maps each comment file to its slide
  ppt/authors.xml                     author names
Anchors come from <ac:spMk id="N"> (a shape id) or are slide-level.

The source file is only ever READ.

    python extract_comments.py <deck.pptx> --round r003
    python extract_comments.py <deck.pptx> --round r003 --out slides/reviews/comment-import-r003.md
    python extract_comments.py <deck.pptx> --dry-run
    python extract_comments.py <deck.pptx> --json
"""
import argparse
import json
import os
import re
import sys
import zipfile

NS_P188 = 'http://schemas.microsoft.com/office/powerpoint/2018/8/main'
NS_A = 'http://schemas.openxmlformats.org/drawingml/2006/main'


def load(deck):
    """-> (comments, authors) where comments is a list of dicts, slide-ordered."""
    from lxml import etree
    comments, authors = [], {}
    with zipfile.ZipFile(deck) as z:
        names = set(z.namelist())

        if 'ppt/authors.xml' in names:
            root = etree.fromstring(z.read('ppt/authors.xml'))
            for a in root.iter('{%s}author' % NS_P188):
                authors[a.get('id')] = a.get('name') or ''

        # comment file -> slide number, via each slide's rels
        cfile_to_slide = {}
        for n in names:
            m = re.match(r'ppt/slides/_rels/slide(\d+)\.xml\.rels$', n)
            if not m:
                continue
            slide_no = int(m.group(1))
            for cf in re.findall(r'modernComment_[0-9A-Fa-f_]+\.xml', z.read(n).decode('utf-8')):
                cfile_to_slide[cf] = slide_no

        for n in sorted(names):
            m = re.match(r'ppt/comments/(modernComment_[0-9A-Fa-f_]+\.xml)$', n)
            if not m:
                continue
            cfile = m.group(1)
            slide_no = cfile_to_slide.get(cfile)
            root = etree.fromstring(z.read(n))
            for cm in root.iter('{%s}cm' % NS_P188):
                sp = cm.find('.//{http://schemas.microsoft.com/office/drawing/2013/main/command}spMk')
                texts = [t.text or '' for t in cm.iter('{%s}t' % NS_A)]
                comments.append({
                    'file': cfile,
                    'slide': slide_no,
                    'anchor_shape_id': int(sp.get('id')) if sp is not None else None,
                    'author': authors.get(cm.get('authorId'), ''),
                    'created': (cm.get('created') or '')[:19],
                    'text': '\n'.join(texts).strip(),
                })
    comments.sort(key=lambda c: (c['slide'] if c['slide'] is not None else 10**6,
                                 c['created']))
    return comments, authors


def anchor_index(deck):
    """{(slide_no, shape_id): text} -- built in one pass, not once per comment."""
    idx = {}
    try:
        from pptx import Presentation
        prs = Presentation(deck)
        for n, slide in enumerate(prs.slides, 1):
            for sh in slide.shapes:
                if sh.has_text_frame:
                    idx[(n, sh.shape_id)] = sh.text_frame.text.replace('\n', ' ')
    except Exception:
        pass
    return idx


def anchor_text(idx, slide_no, shape_id, width=60):
    """Best-effort: the text of the shape a comment is anchored to."""
    if slide_no is None or shape_id is None:
        return ''
    return idx.get((slide_no, shape_id), '')[:width]


def render_md(deck, comments, authors, rnd):
    import datetime
    idx = anchor_index(deck)
    slides = sorted({c['slide'] for c in comments if c['slide'] is not None})
    who = ' / '.join(sorted({c['author'] for c in comments if c['author']})) or '(unknown)'
    L = []
    L.append('# Comment Import — %s\n' % rnd)
    L.append('## Source\n')
    L.append('* レビュー入力: `%s`' % deck.replace('\\', '/'))
    L.append('* 入力モード: <A: 別ファイル / B: 作業中 deck に直接付与 → 編集前にスナップショット>')
    L.append('* 抽出日: %s' % datetime.date.today().isoformat())
    L.append('* 作成者: %s' % who)
    L.append('* コメント形式: PowerPoint modern comments（`ppt/comments/modernComment_*.xml`）')
    L.append('* 抽出件数: **%d件 / %dスライド**（%s）\n'
             % (len(comments), len(slides),
                ', '.join(str(s) for s in slides) if slides else '-'))
    L.append('## Round の性格\n')
    L.append('<!-- 前ラウンドへの再フィードバックか、新規章の初レビューか等 -->\n')
    L.append('---\n')
    L.append('## コメント\n')
    for i, c in enumerate(comments, 1):
        cid = '%s-%03d' % (rnd.upper().replace('R', 'CMT-r'), i) \
            if rnd.lower().startswith('r') else 'CMT-%s-%03d' % (rnd, i)
        if c['anchor_shape_id'] is not None:
            at = anchor_text(idx, c['slide'], c['anchor_shape_id'])
            anchor = 'sp=%s%s' % (c['anchor_shape_id'], '（%s）' % at if at else '')
        else:
            anchor = '`slide-level`'
        L.append('### %s — Slide %s / %s / `[分類][重要度]`\n'
                 % (cid, c['slide'] if c['slide'] is not None else '?', anchor))
        L.append('原文:\n')
        L.append('```text')
        L.append(c['text'])
        L.append('```\n')
        L.append('対応 `OPEN`: \n')
        L.append('---\n')
    L.append('## 技術確認の記録\n')
    L.append('<!-- 一次情報で確認した内容、そのまま反映しなかった判断とその根拠 -->\n')
    L.append('## 波及・要確認\n')
    L.append('<!-- 他スライドへの影響（変更はせず報告のみ）／ユーザーに再確認したい判断 -->\n')
    L.append('## Not applied / Rejected\n')
    L.append('<!-- REJECTED のコメントと根拠。無ければ「なし」と明記 -->\n')
    return '\n'.join(L) + '\n'


def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('deck')
    ap.add_argument('--round', dest='rnd', default='rNNN', help='review round, e.g. r003')
    ap.add_argument('--out', help='write markdown here (default: stdout)')
    ap.add_argument('--dry-run', action='store_true', help='summary only, no markdown')
    ap.add_argument('--json', action='store_true')
    a = ap.parse_args()

    sys.stdout.reconfigure(encoding='utf-8')
    comments, authors = load(a.deck)

    if a.json:
        print(json.dumps(comments, ensure_ascii=False, indent=2))
        return
    if a.dry_run:
        slides = sorted({c['slide'] for c in comments if c['slide'] is not None})
        print('%d comments on %d slides: %s'
              % (len(comments), len(slides), ', '.join(str(s) for s in slides)))
        for c in comments:
            anc = ('sp=%s' % c['anchor_shape_id']) if c['anchor_shape_id'] else 'slide-level'
            first = c['text'].splitlines()[0] if c['text'] else ''
            print('  slide %-3s %-13s %-19s %s'
                  % (c['slide'], anc, c['created'], first[:58]))
        return

    md = render_md(a.deck, comments, authors, a.rnd)
    if a.out:
        os.makedirs(os.path.dirname(a.out) or '.', exist_ok=True)
        if os.path.exists(a.out):
            sys.exit('refusing to overwrite existing %s' % a.out)
        with open(a.out, 'w', encoding='utf-8') as f:
            f.write(md)
        print('wrote %s (%d comments)' % (a.out, len(comments)))
    else:
        print(md)


if __name__ == '__main__':
    main()
