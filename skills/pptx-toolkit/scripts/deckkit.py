# -*- coding: utf-8 -*-
"""Primitives for creating and editing a .pptx with python-pptx + raw OOXML.

Import this from a throwaway batch script; do not run it directly.

    import sys
    sys.path.insert(0, r'<pptx-toolkit>/scripts')
    from pptx import Presentation
    from deckkit import R, set_runs, shape_by_id, rrect, textbox, connector

Four rules this module exists to enforce (each cost real debugging time):

1. Build XML with ``pptx.oxml.parse_xml``, never ``lxml.etree.fromstring``. Plain lxml
   elements lack python-pptx's custom element classes, and ``slide.shapes`` iteration
   then dies with ``'_Element' object has no attribute 'has_ph_elm'``. ``frag()`` is correct.
2. When deleting shapes by ``p:cNvPr/@id``, only touch ``sp|cxnSp|pic|graphicFrame|grpSp``.
   ``<p:nvGrpSpPr>`` also carries a ``cNvPr`` (id="1"); removing it corrupts the spTree and
   PowerPoint then refuses the file ("ファイルを開けません") while python-pptx still opens
   it happily. ``delete_ids()`` filters correctly.
3. ``<a:endParaRPr>`` must stay the LAST child of its paragraph. Append runs after a
   leftover endParaRPr and the paragraph renders BLANK in PowerPoint even though the XML
   looks right and python-pptx reads it back fine. ``set_para()`` handles this.
4. Deleting slides needs the drop_rel dance -- see ``delete_slides()``.

Always ``save_backup()`` before a batch edit. Batch edits are not undoable.
"""
import copy
import shutil
import time
from lxml import etree

NS = {
    'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
    'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
}
A = NS['a']
P = NS['p']
EMU = 914400


def emu(inches):
    return int(round(inches * EMU))


def _q(tag):
    ns, _, local = tag.partition(':')
    return '{%s}%s' % (NS[ns], local)


def X(tag, **attrs):
    el = etree.SubElement(etree.Element('root'), _q(tag))
    for k, v in attrs.items():
        el.set(k, str(v))
    return el


def frag(xml):
    """Parse an XML fragment through python-pptx's oxml parser.

    Using parse_xml (not plain lxml) is required so the resulting elements get
    python-pptx's custom element classes; otherwise slide.shapes iteration
    blows up on them.
    """
    from pptx.oxml import parse_xml
    wrapper = ('<root xmlns:a="%s" xmlns:p="%s" '
               'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">%s</root>'
               % (A, P, xml))
    return list(parse_xml(wrapper))


def spTree(slide):
    return slide.shapes._spTree


def next_id(slide):
    ids = [int(e.get('id')) for e in spTree(slide).iter(_q('p:cNvPr'))]
    return max(ids) + 1


def shape_by_id(slide, sid):
    for sh in slide.shapes:
        if sh.shape_id == sid:
            return sh
    raise KeyError('shape id %s not found' % sid)


# Only these spTree children are real shapes. <p:nvGrpSpPr> also carries a cNvPr
# (id="1") -- deleting it corrupts the spTree and PowerPoint then refuses the file.
SHAPE_TAGS = frozenset(('sp', 'cxnSp', 'pic', 'graphicFrame', 'grpSp'))


def is_shape_el(el):
    return etree.QName(el).localname in SHAPE_TAGS


def shape_elements(slide):
    """The spTree children that are actual shapes (never nvGrpSpPr / grpSpPr)."""
    return [el for el in spTree(slide) if is_shape_el(el)]


def el_by_id(slide, sid):
    for el in shape_elements(slide):
        cnv = el.find('.//' + _q('p:cNvPr'))
        if cnv is not None and cnv.get('id') == str(sid):
            return el
    raise KeyError('element id %s not found' % sid)


def delete_ids(slide, ids):
    for sid in ids:
        el = el_by_id(slide, sid)
        el.getparent().remove(el)


def delete_range(slide, lo, hi, keep=()):
    """Delete every shape whose cNvPr id is in [lo, hi], except ids in `keep`.

    Safe against the nvGrpSpPr trap: non-shape elements are skipped outright.
    Useful when rebuilding a whole diagram while keeping title/footer shapes.
    """
    keep = {str(k) for k in keep}
    removed = []
    for el in list(shape_elements(slide)):
        cnv = el.find('.//' + _q('p:cNvPr'))
        if cnv is None:
            continue
        sid = cnv.get('id')
        if sid in keep:
            continue
        if lo <= int(sid) <= hi:
            el.getparent().remove(el)
            removed.append(int(sid))
    return sorted(removed)


def move_after(slide, el, ref_el):
    """Reorder el to sit immediately after ref_el in z-order."""
    tree = spTree(slide)
    tree.remove(el)
    ref_el.addnext(el)


def move_before(slide, el, ref_el):
    tree = spTree(slide)
    tree.remove(el)
    ref_el.addprevious(el)


def set_xfrm(el, x=None, y=None, w=None, h=None):
    xfrm = el.find('.//' + _q('a:xfrm'))
    off = xfrm.find(_q('a:off'))
    ext = xfrm.find(_q('a:ext'))
    if x is not None:
        off.set('x', str(emu(x)))
    if y is not None:
        off.set('y', str(emu(y)))
    if w is not None:
        ext.set('cx', str(emu(w)))
    if h is not None:
        ext.set('cy', str(emu(h)))


# --------------------------------------------------------------------------
# run / paragraph helpers
# --------------------------------------------------------------------------

# Fallback colours used when R()/connector() omit `color`. These are DARK-theme
# values; on a light deck they would be invisible. Call use_defaults() from the
# profile at the top of a batch script -- see the light-preset note in SKILL.md.
DEFAULTS = {'text': 'ECEAF7', 'rule': '3C3463', 'font': 'Calibri'}


def use_defaults(text=None, rule=None, font=None):
    """Point the R()/connector() fallbacks at this project's palette.

        prof = load_profile('slides/deck-profile.json'); pal = prof['palette']
        use_defaults(text=pal['text'], rule=pal['rule'])

    Only affects calls that OMIT `color`; explicit colours are untouched.
    """
    if text:
        DEFAULTS['text'] = text
    if rule:
        DEFAULTS['rule'] = rule
    if font:
        DEFAULTS['font'] = font
    return dict(DEFAULTS)


def R(text, sz=9.0, color=None, font=None, b=False, i=False, hl=None, u=False,
      a=None):
    """Describe one text run. color/font default to DEFAULTS (see use_defaults).

    hl paints a highlighter behind the run (hex, no '#'); u underlines it. Both
    exist so a key word can be marked WITHOUT splitting it into its own shape --
    emphasis inside a sentence is the whole point of a key message.

    `a` is the ink's opacity, 0.0-1.0 (None = fully opaque). PowerPoint honours
    alpha on run text, so a span can be pushed BACK without changing its colour.
    That is how codehl's focus mode dims the lines around the marked one while
    keeping the syntax colours -- and it keeps check_deck's palette check quiet,
    because the colour value written to the file is still a profile colour.
    """
    return dict(text=text, sz=sz, color=color or DEFAULTS['text'],
                font=font or DEFAULTS['font'], b=b, i=i, hl=hl, u=u, a=a)


def _esc(s):
    return (s.replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;'))


def _run_xml(r):
    # a:rPr child order is fixed by the schema: solidFill, then highlight, then
    # latin/ea/cs. Emitting highlight after latin makes PowerPoint reject the run.
    hl = r.get('hl')
    hl_xml = '<a:highlight><a:srgbClr val="%s"/></a:highlight>' % hl if hl else ''
    u_attr = ' u="sng"' if r.get('u') else ''
    a = r.get('a')
    if a is None:
        clr_xml = '<a:srgbClr val="%s"/>' % r['color']
    else:
        clr_xml = ('<a:srgbClr val="%s"><a:alpha val="%d"/></a:srgbClr>'
                   % (r['color'],
                      max(0, min(100000, int(round(float(a) * 100000))))))
    return (
        '<a:r><a:rPr sz="%d" b="%d" i="%d"%s><a:solidFill>%s</a:solidFill>'
        '%s<a:latin typeface="%s"/><a:ea typeface="%s"/><a:cs typeface="%s"/></a:rPr>'
        '<a:t>%s</a:t></a:r>'
        % (int(round(r['sz'] * 100)), 1 if r['b'] else 0, 1 if r['i'] else 0,
           u_attr, clr_xml, hl_xml, r['font'], r['font'], r['font'],
           _esc(r['text']))
    )


def _para_xml(runs, algn):
    body = ''.join(_run_xml(r) for r in runs)
    return '<a:p><a:pPr algn="%s"><a:buNone/></a:pPr>%s</a:p>' % (algn, body)


def textbox(slide, x, y, w, h, paras, algn='l', anchor='ctr', wrap='square',
            spc_pct=None):
    """paras: list of list-of-R (one inner list per paragraph)."""
    sid = next_id(slide)
    lnspc = ''
    if spc_pct:
        lnspc = '<a:lnSpc><a:spcPct val="%d"/></a:lnSpc>' % int(spc_pct * 1000)
    ps = []
    for runs in paras:
        body = ''.join(_run_xml(r) for r in runs)
        ps.append('<a:p><a:pPr algn="%s">%s<a:buNone/></a:pPr>%s</a:p>' % (algn, lnspc, body))
    xml = (
        '<p:sp><p:nvSpPr><p:cNvPr id="%d" name="TextBox %d"/><p:cNvSpPr txBox="1"/><p:nvPr/></p:nvSpPr>'
        '<p:spPr><a:xfrm><a:off x="%d" y="%d"/><a:ext cx="%d" cy="%d"/></a:xfrm>'
        '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom><a:noFill/></p:spPr>'
        '<p:txBody><a:bodyPr wrap="%s" lIns="0" tIns="0" rIns="0" bIns="0" anchor="%s"/>'
        '<a:lstStyle/>%s</p:txBody></p:sp>'
        % (sid, sid - 1, emu(x), emu(y), emu(w), emu(h), wrap, anchor, ''.join(ps))
    )
    el = frag(xml)[0]
    spTree(slide).append(el)
    return el


_STYLE = ('<p:style><a:lnRef idx="1"><a:schemeClr val="accent1"/></a:lnRef>'
          '<a:fillRef idx="3"><a:schemeClr val="accent1"/></a:fillRef>'
          '<a:effectRef idx="2"><a:schemeClr val="accent1"/></a:effectRef>'
          '<a:fontRef idx="minor"><a:schemeClr val="lt1"/></a:fontRef></p:style>')

_SHADOW = ('<a:effectLst><a:outerShdw blurRad="101600" dist="38100" dir="5400000" '
           'algn="bl" rotWithShape="0"><a:srgbClr val="000000"><a:alpha val="22000"/>'
           '</a:srgbClr></a:outerShdw></a:effectLst>')


def rrect(slide, x, y, w, h, fill=None, line=None, lw=15240, adj=14000,
          shadow=False, fill_alpha=None, dash=None, runs=None, **kw):
    """Rounded rectangle. A thin wrapper over shape() kept for compatibility --
    every existing batch script calls this. New code can call shape() directly
    with any prstGeom."""
    kw.setdefault('name', 'Rounded Rectangle')
    return shape(slide, 'roundRect', x, y, w, h, fill=fill, line=line, lw=lw,
                 adj=adj, shadow=shadow, fill_alpha=fill_alpha, dash=dash,
                 runs=runs, **kw)


def connector(slide, x, y, w, h, color=None, lw=9525, dash=None,
              tail=None, head=None, flipH=False, flipV=False,
              start=None, end=None):
    """Straight line, positioned by BOUNDING BOX -- not by two points.

    Arrowheads: read gotcha (11) before reaching for head=/tail=.

        head=   writes <a:headEnd>, and OOXML puts that at the line's START
        tail=   writes <a:tailEnd>, and OOXML puts that at the line's END

    So `head` does NOT mean "the pointy end the arrow points at" -- it is the
    end the arrow comes FROM. Both names are kept as literal OOXML pass-throughs
    so old batch scripts keep rendering identically. New code should use the
    unambiguous aliases:

        start=  decoration at the line's start point   (same wire format as head=)
        end=    decoration at the line's end point     (same wire format as tail=)

    With no flips the start point is the box corner (x, y) and the end point is
    (x+w, y+h). flipH/flipV mirror the box, and therefore swap which physical
    corner is the start.

    For "an arrow from A to B", do not work the box out by hand -- call
    arrow(slide, ax, ay, bx, by), which derives box and flips for you and lands
    the arrowhead on B every time.
    """
    if start is not None:
        if head is not None:
            raise ValueError('connector(): pass start= or head=, not both')
        head = start
    if end is not None:
        if tail is not None:
            raise ValueError('connector(): pass end= or tail=, not both')
        tail = end
    color = color or DEFAULTS['rule']
    sid = next_id(slide)
    dash_xml = '<a:prstDash val="%s"/>' % dash if dash else ''
    tail_xml = '<a:tailEnd type="%s" w="med" len="med"/>' % tail if tail else ''
    head_xml = '<a:headEnd type="%s" w="med" len="med"/>' % head if head else ''
    flips = ''
    if flipH:
        flips += ' flipH="1"'
    if flipV:
        flips += ' flipV="1"'
    xml = (
        '<p:cxnSp><p:nvCxnSpPr><p:cNvPr id="%d" name="Connector %d"/><p:cNvCxnSpPr/><p:nvPr/></p:nvCxnSpPr>'
        '<p:spPr><a:xfrm%s><a:off x="%d" y="%d"/><a:ext cx="%d" cy="%d"/></a:xfrm>'
        '<a:prstGeom prst="line"><a:avLst/></a:prstGeom>'
        '<a:ln w="%d"><a:solidFill><a:srgbClr val="%s"/></a:solidFill>%s%s%s</a:ln></p:spPr>'
        '<p:style><a:lnRef idx="2"><a:schemeClr val="accent1"/></a:lnRef>'
        '<a:fillRef idx="0"><a:schemeClr val="accent1"/></a:fillRef>'
        '<a:effectRef idx="1"><a:schemeClr val="accent1"/></a:effectRef>'
        '<a:fontRef idx="minor"><a:schemeClr val="tx1"/></a:fontRef></p:style></p:cxnSp>'
        % (sid, sid - 1, flips, emu(x), emu(y), emu(w), emu(h), lw, color,
           dash_xml, head_xml, tail_xml)
    )
    el = frag(xml)[0]
    spTree(slide).append(el)
    return el


def arrow(slide, x1, y1, x2, y2, color=None, lw=9525, dash=None,
          tip='triangle', both=False, back=None):
    """Arrow from (x1, y1) to (x2, y2). The arrowhead lands on (x2, y2) -- always.

    Use this for every conceptual diagram. It derives the bounding box and the
    flipH/flipV that connector() wants, so the caller never has to reason about
    which box corner OOXML considers the line's start (gotcha 11) -- the single
    biggest source of backwards arrows in a finished deck.

        arrow(s, 3.0, 2.0, 6.5, 2.0)              # --> head at (6.5, 2.0)
        arrow(s, 6.5, 2.0, 3.0, 2.0)              # <-- head at (3.0, 2.0)
        arrow(s, 4.0, 5.0, 4.0, 2.0)              # upward, head at the top
        arrow(s, 3.0, 2.0, 6.5, 4.0, both=True)   # <-> heads on both ends
        arrow(s, ..., tip=None)                   # plain line, no head

    Write the coordinates in the direction the *meaning* flows: for "B社 supplies
    workers to C社", pass B's edge first and C's edge second, and the picture
    cannot disagree with the caption.

    tip   arrowhead style at (x2, y2): 'triangle' (default), 'stealth', 'arrow',
          'oval', 'diamond', or None for no head.
    both  also put a `tip`-styled head on (x1, y1).
    back  arrowhead style at (x1, y1); overrides `both` when given.
    """
    x, y = min(x1, x2), min(y1, y2)
    w, h = abs(x2 - x1), abs(y2 - y1)
    # An unflipped line runs from the box's top-left to its bottom-right, so
    # mirror the box on whichever axis the drawing direction disagrees with.
    if back is None and both:
        back = tip
    return connector(slide, x, y, w, h, color=color, lw=lw, dash=dash,
                     start=back, end=tip, flipH=x2 < x1, flipV=y2 < y1)


def set_para(shape, pidx, runs):
    """Replace the runs of a single paragraph, preserving pPr and keeping any
    a:endParaRPr as the LAST child (OOXML requires endParaRPr after all runs)."""
    p_el = shape.text_frame.paragraphs[pidx]._p
    for r in p_el.findall(_q('a:r')):
        p_el.remove(r)
    epr = p_el.find(_q('a:endParaRPr'))
    for rr in runs:
        node = frag(_run_xml(rr))[0]
        if epr is not None:
            epr.addprevious(node)
        else:
            p_el.append(node)


def set_runs(shape, paras):
    """Replace a text frame's contents, keeping the first paragraph's pPr."""
    txBody = shape.text_frame._txBody
    keep_pPr = None
    for p_el in txBody.findall(_q('a:p')):
        if keep_pPr is None:
            pPr = p_el.find(_q('a:pPr'))
            if pPr is not None:
                keep_pPr = copy.deepcopy(pPr)
        txBody.remove(p_el)
    for runs in paras:
        p_el = etree.SubElement(txBody, _q('a:p'))
        if keep_pPr is not None:
            p_el.append(copy.deepcopy(keep_pPr))
        for r in runs:
            p_el.append(frag(_run_xml(r))[0])


# --------------------------------------------------------------------------
# slide-level operations
# --------------------------------------------------------------------------

def save_backup(path, suffix=None):
    """Copy a .pptx next to itself with a timestamp. Call BEFORE any batch edit.

    Returns the backup path. Batch edits are not undoable; this is the only net.
    """
    suffix = suffix or time.strftime('%Y%m%d-%H%M%S')
    dst = '%s.bak-%s.pptx' % (path[:-5] if path.endswith('.pptx') else path, suffix)
    shutil.copy2(path, dst)
    return dst


def delete_slides(prs, positions):
    """Delete slides by 1-based position. Handles the drop_rel dance correctly.

    Deletes highest position first so earlier indices stay valid. Comments and
    notes anchored to a deleted slide drop with it.

    Returns the list of deleted positions.
    """
    from pptx.oxml.ns import qn
    sldIdLst = prs.slides._sldIdLst
    ids = list(sldIdLst)
    for pos in sorted(set(positions), reverse=True):
        if not (1 <= pos <= len(ids)):
            raise IndexError('slide position %s out of range 1..%s' % (pos, len(ids)))
        sldId = ids[pos - 1]
        prs.part.drop_rel(sldId.get(qn('r:id')))
        sldIdLst.remove(sldId)
    return sorted(set(positions), reverse=True)


def move_slide(prs, frm, to):
    """Move the slide at 1-based position `frm` to 1-based position `to`."""
    sldIdLst = prs.slides._sldIdLst
    ids = list(sldIdLst)
    n = len(ids)
    if not (1 <= frm <= n and 1 <= to <= n):
        raise IndexError('positions must be within 1..%s' % n)
    el = ids[frm - 1]
    sldIdLst.remove(el)
    ref = list(sldIdLst)
    if to - 1 >= len(ref):
        sldIdLst.append(el)
    else:
        ref[to - 1].addprevious(el)


def slide_titles(prs, limit=None):
    """[(position, first non-empty text)] -- for confirming what you are about to delete."""
    out = []
    for i, s in enumerate(prs.slides, 1):
        text = ''
        for sh in s.shapes:
            if sh.has_text_frame and sh.text_frame.text.strip():
                text = sh.text_frame.text.strip().splitlines()[0]
                break
        out.append((i, text[:limit] if limit else text))
    return out


# --------------------------------------------------------------------------
# deck creation
#
# The look of a deck is NOT carried by its theme. This toolchain writes every
# fill/line/run as an explicit srgbClr, and paints the canvas with an explicit
# <p:bg>. The theme is decorative -- see set_theme_colors().
# --------------------------------------------------------------------------

def load_profile(path):
    """Read a deck-profile.json. Returns the dict (palette under ['palette'])."""
    import json
    with open(path, encoding='utf-8') as f:
        return json.load(f)


def set_deck_size(prs, width_in=10.0, height_in=5.625, sz_type='screen16x9'):
    """Set slide dimensions AND the <p:sldSz @type> that python-pptx leaves stale.

    Assigning prs.slide_width/height alone leaves type="screen4x3" on a 16:9
    deck, so PowerPoint reports the wrong preset under Design > Slide Size.
    Pass sz_type=None to drop the attribute (correct for odd ratios).
    """
    prs.slide_width = emu(width_in)
    prs.slide_height = emu(height_in)
    sz = prs._element.find(_q('p:sldSz'))
    if sz is not None:
        if sz_type:
            sz.set('type', sz_type)
        elif 'type' in sz.attrib:
            del sz.attrib['type']
    return prs


def new_presentation(width_in=10.0, height_in=5.625, sz_type='screen16x9'):
    """A zero-slide Presentation at the given size, from python-pptx's template."""
    from pptx import Presentation
    return set_deck_size(Presentation(), width_in, height_in, sz_type)


def blank_layout(prs):
    """The first layout with no cloneable placeholders.

    Building a slide on a layout that HAS placeholders inherits the master's
    txStyles (44pt title / 32pt bulleted body, theme text colour), which is
    wrong for a free-shape deck and unreadable on a dark canvas.
    """
    preferred = ('DEFAULT', 'Blank', '空白')
    cands = [l for l in prs.slide_layouts
             if not list(l.iter_cloneable_placeholders())]
    if not cands:
        raise ValueError('no placeholder-free layout in this deck')
    for name in preferred:
        for l in cands:
            if l.name == name:
                return l
    return cands[0]


def _cSld(obj):
    """The <p:cSld> of a slide, layout or master."""
    el = getattr(obj, '_element', None)
    if el is None:
        raise TypeError('expected a slide, layout or master')
    c = el.find(_q('p:cSld'))
    if c is None:
        raise ValueError('no <p:cSld> on this part')
    return c


def set_bg(obj, color):
    """Solid background on a slide, layout or master.

    <p:bg> must be the FIRST child of <p:cSld> (schema order: bg?, spTree, ...).
    Replaces any existing <p:bg>. This is how the canvas colour is actually
    carried -- the theme/master chain otherwise resolves to white.
    """
    c = _cSld(obj)
    for old in c.findall(_q('p:bg')):
        c.remove(old)
    bg = frag('<p:bg><p:bgPr><a:solidFill><a:srgbClr val="%s"/></a:solidFill>'
              '<a:effectLst/></p:bgPr></p:bg>' % color)[0]
    c.insert(0, bg)
    return bg


def clear_bg(obj):
    """Remove <p:bg> so this part inherits its parent's background."""
    c = _cSld(obj)
    n = 0
    for old in c.findall(_q('p:bg')):
        c.remove(old)
        n += 1
    return n


def add_slide(prs, layout=None, bg=None):
    """Append an empty slide and return it, ready for textbox()/rrect()/connector().

    layout defaults to blank_layout(prs). bg, when given, paints this slide's
    canvas; usually unnecessary because new_deck.py backgrounds the master, but
    it matches how real decks in the wild carry it (per-slide).
    """
    slide = prs.slides.add_slide(layout or blank_layout(prs))
    if bg:
        set_bg(slide, bg)
    return slide


def apply_bg_all(prs, color):
    """set_bg on every slide. For retrofitting a deck that lacks a background."""
    for slide in prs.slides:
        set_bg(slide, color)
    return len(prs.slides._sldIdLst)


def set_theme_colors(prs, colors):
    """Patch <a:clrScheme> accents in ppt/theme/theme1.xml.

    colors: {'accent1': 'FFB454', ..., 'hlink': '...', 'folHlink': '...'}

    Only accent1..6/hlink/folHlink are honoured. dk1/lt1/dk2/lt2 and the
    master's <p:clrMap> drive PowerPoint's light/dark contrast heuristics --
    changing them makes tables and SmartArt pick unreadable text colours.

    Nothing this toolchain draws reads the theme (every colour is an explicit
    srgbClr), so this only matters for content PowerPoint styles itself:
    tables, charts, SmartArt, and the colour picker a human sees.

    python-pptx does not model the theme part, so this edits Part._blob.
    """
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    allowed = ('accent1', 'accent2', 'accent3', 'accent4', 'accent5',
               'accent6', 'hlink', 'folHlink')
    wanted = {k: v for k, v in (colors or {}).items() if k in allowed}
    if not wanted:
        return []
    theme = prs.slide_masters[0].part.part_related_by(RT.THEME)
    root = etree.fromstring(theme._blob)
    scheme = root.find('.//' + _q('a:clrScheme'))
    if scheme is None:
        return []
    done = []
    for slot, val in wanted.items():
        node = scheme.find(_q('a:' + slot))
        if node is None:
            continue
        for child in list(node):
            node.remove(child)
        node.append(frag('<a:srgbClr val="%s"/>' % val)[0])
        done.append(slot)
    theme._blob = etree.tostring(root, xml_declaration=True,
                                 encoding='UTF-8', standalone=True)
    return sorted(done)


# --------------------------------------------------------------------------
# generic shapes
#
# rrect() used to be the ONLY filled shape this module could draw, which is why
# decks built with it look uniform: every emphasis had to be another rounded box
# in the same fill. shape() opens the whole prstGeom catalogue with one call, so
# an icon, a UML decision diamond, a chevron step or a callout is one line.
# --------------------------------------------------------------------------

def _avlst(adj):
    """<a:avLst> from adj.

    None -> empty; int/float -> name="adj"; list -> adj1, adj2, ...;
    dict -> the given names (chevron uses "adj", wedgeRectCallout adj1..adj3).
    """
    if adj is None:
        return '<a:avLst/>'
    if isinstance(adj, dict):
        items = list(adj.items())
    elif isinstance(adj, (list, tuple)):
        items = [('adj%d' % (i + 1), v) for i, v in enumerate(adj)]
    else:
        items = [('adj', adj)]
    return '<a:avLst>%s</a:avLst>' % ''.join(
        '<a:gd name="%s" fmla="val %d"/>' % (k, int(v)) for k, v in items)


def _fill_xml(fill, fill_alpha):
    if fill is None:
        return '<a:noFill/>'
    if fill_alpha is not None:
        return ('<a:solidFill><a:srgbClr val="%s"><a:alpha val="%d"/></a:srgbClr>'
                '</a:solidFill>' % (fill, int(fill_alpha * 1000)))
    return '<a:solidFill><a:srgbClr val="%s"/></a:solidFill>' % fill


def _line_xml(line, lw, dash, line_alpha=None, cap=None):
    if line is None:
        return '<a:ln><a:noFill/></a:ln>'
    alpha = ('<a:alpha val="%d"/>' % int(line_alpha * 1000)) if line_alpha is not None else ''
    dash_xml = '<a:prstDash val="%s"/>' % dash if dash else ''
    cap_attr = ' cap="%s"' % cap if cap else ''
    return ('<a:ln w="%d"%s><a:solidFill><a:srgbClr val="%s">%s</a:srgbClr></a:solidFill>%s</a:ln>'
            % (lw, cap_attr, line, alpha, dash_xml))


def _norm_paras(paras):
    """Accept a str, one list-of-R, or a list of list-of-R -> list of list-of-R."""
    if paras is None:
        return None
    if isinstance(paras, str):
        return [[R(paras)]]
    paras = list(paras)
    if paras and isinstance(paras[0], dict):      # a single paragraph's runs
        return [paras]
    return paras


def _ins_attrs(ins):
    """bodyPr inset attributes. ins = None (PowerPoint default) | inches | (l,t,r,b)."""
    if ins is None:
        return ''
    if isinstance(ins, (list, tuple)):
        l, t, r, b = ins
    else:
        l = t = r = b = ins
    return ' lIns="%d" tIns="%d" rIns="%d" bIns="%d"' % (emu(l), emu(t), emu(r), emu(b))


def _txbody_xml(paras, algn, anchor, wrap, spc_pct, ins, autofit=None):
    lnspc = '<a:lnSpc><a:spcPct val="%d"/></a:lnSpc>' % int(spc_pct * 1000) if spc_pct else ''
    fit = ''
    if autofit == 'shrink':
        fit = '<a:normAutofit/>'
    elif autofit == 'resize':
        fit = '<a:spAutoFit/>'
    if not paras:
        ps = '<a:p><a:pPr algn="%s"/><a:endParaRPr/></a:p>' % algn
    else:
        ps = ''.join(
            '<a:p><a:pPr algn="%s">%s<a:buNone/></a:pPr>%s</a:p>'
            % (algn, lnspc, ''.join(_run_xml(r) for r in runs))
            for runs in paras)
    return ('<p:txBody><a:bodyPr rtlCol="0" wrap="%s" anchor="%s"%s>%s</a:bodyPr>'
            '<a:lstStyle/>%s</p:txBody>' % (wrap, anchor, _ins_attrs(ins), fit, ps))


def shape(slide, prst, x, y, w, h, fill=None, line=None, lw=15240, adj=None,
          shadow=False, fill_alpha=None, dash=None, line_alpha=None,
          rot=None, flipH=False, flipV=False,
          runs=None, algn='ctr', anchor='ctr', wrap='square', spc_pct=None,
          ins=None, name=None, style=True):
    """Draw any PowerPoint preset geometry, optionally with its own text.

    prst is any ST_ShapeType value -- 'ellipse', 'triangle', 'chevron', 'can',
    'donut', 'rightArrow', 'lightningBolt', 'noSmoking', 'flowChartDecision',
    'flowChartDocument', 'wedgeRectCallout', 'star5', 'mathMultiply', ...

    `runs` puts text INSIDE the shape, so a labelled box is one shape, not two.
    That matters: frame-and-label-as-separate-shapes is the single most common
    source of "I moved the box but the text stayed behind".

    rot is in degrees (converted to the 60000ths OOXML wants).
    """
    sid = next_id(slide)
    label = name or ''.join(
        (' ' + c if c.isupper() else c) for c in prst).strip().title()
    rot_attr = ' rot="%d"' % int(round(rot * 60000)) if rot else ''
    if flipH:
        rot_attr += ' flipH="1"'
    if flipV:
        rot_attr += ' flipV="1"'
    xml = (
        '<p:sp><p:nvSpPr><p:cNvPr id="%d" name="%s %d"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>'
        '<p:spPr><a:xfrm%s><a:off x="%d" y="%d"/><a:ext cx="%d" cy="%d"/></a:xfrm>'
        '<a:prstGeom prst="%s">%s</a:prstGeom>%s%s%s</p:spPr>%s%s</p:sp>'
        % (sid, label, sid - 1, rot_attr, emu(x), emu(y), emu(w), emu(h),
           prst, _avlst(adj), _fill_xml(fill, fill_alpha),
           _line_xml(line, lw, dash, line_alpha),
           _SHADOW if shadow else '<a:effectLst/>',
           _STYLE if style else '',
           _txbody_xml(_norm_paras(runs), algn, anchor, wrap, spc_pct, ins))
    )
    el = frag(xml)[0]
    spTree(slide).append(el)
    return el


def text_in(el, paras, algn='ctr', anchor='ctr', wrap='square', spc_pct=None,
            ins=0.06, autofit=None):
    """Put text inside an EXISTING shape element, replacing its txBody.

    Accepts a raw lxml element (what shape()/rrect() return) or a python-pptx
    shape. Use this instead of stacking a separate textbox on top of a box.
    """
    sp = getattr(el, '_element', el)
    body = sp.find(_q('p:txBody'))
    new = frag(_txbody_xml(_norm_paras(paras), algn, anchor, wrap, spc_pct,
                           ins, autofit))[0]
    if body is not None:
        sp.replace(body, new)
    else:
        sp.append(new)
    return el


def picture(slide, path, x, y, w=None, h=None):
    """Insert an image. Give w or h alone and the aspect ratio fills the other.

    Returns the raw <p:pic> element, so set_xfrm()/move_before() work on it.
    """
    from pptx.util import Inches
    pic = slide.shapes.add_picture(
        path, Inches(x), Inches(y),
        Inches(w) if w is not None else None,
        Inches(h) if h is not None else None)
    return pic._element


# --------------------------------------------------------------------------
# text measurement and Japanese line breaking
#
# ADVISORY ONLY. Proportional metrics vary by font; this exists to catch a line
# that will obviously wrap BEFORE you spend a render cycle finding out.
# It never replaces looking at the PNG.
# --------------------------------------------------------------------------

MONO_FONTS = frozenset((
    'Courier New', 'Consolas', 'Menlo', 'Monaco', 'MS Gothic', 'BIZ UDGothic',
    'Cascadia Code', 'Cascadia Mono', 'SF Mono', 'DejaVu Sans Mono',
))

# Ranges that render full-width (1 em) in a CJK font.
_FULLWIDTH_RANGES = (
    (0x1100, 0x115F), (0x2E80, 0x303E), (0x3041, 0x33FF), (0x3400, 0x4DBF),
    (0x4E00, 0x9FFF), (0xA000, 0xA4CF), (0xAC00, 0xD7A3), (0xF900, 0xFAFF),
    (0xFE30, 0xFE4F), (0xFF00, 0xFF60), (0xFFE0, 0xFFE6),
    (0x1F300, 0x1FAFF),
)

_NARROW = set(chr(0x69) + chr(0x6C) + chr(0x6A) + chr(0x74) + chr(0x66) +
              chr(0x49) + '.,:;|!()[]{}-' + chr(0x27) + chr(0x22) + chr(0x60))
_WIDE = set('MWmw@%&')


def is_fullwidth(ch):
    o = ord(ch)
    return any(lo <= o <= hi for lo, hi in _FULLWIDTH_RANGES)


def _em(ch, mono):
    """Width of one character in em units. Rough but stable."""
    if ch == '\t':
        return 4 * (0.6 if mono else 0.26)
    if ch == ' ':
        return 0.6 if mono else 0.26
    if is_fullwidth(ch):
        return 1.0
    if mono:
        return 0.6
    if ch in _NARROW:
        return 0.30
    if ch in _WIDE:
        return 0.86
    if ch.isdigit() or ch.isupper():
        return 0.60
    return 0.50


def measure(runs, font=None, sz=None):
    """Estimated drawn width in INCHES of one paragraph.

    runs: a list of R() dicts, or a plain string (then pass font= and sz=).
    """
    if isinstance(runs, str):
        runs = [R(runs, sz or 9.0, font=font)]
    elif isinstance(runs, dict):          # a single R(), not a list of them
        runs = [runs]
    total = 0.0
    for r in runs:
        unit = r['sz'] / 72.0
        mono = (r.get('font') or '') in MONO_FONTS
        for ch in r['text']:
            total += unit * _em(ch, mono)
    return total


def fits(runs, width_in, font=None, sz=None, slack=0.98):
    """True when the paragraph is expected to stay on one line inside width_in."""
    return measure(runs, font, sz) <= width_in * slack


# 行頭禁則: these may not START a line.
NO_LINE_START = (
    '。、，．・：；！？'
    '）」』】〉》〕｝］”’'
    '!?,.:;)]}>'
    'ーヽヾゝゞ々‐–—～'
    'ぁぃぅぇぉっゃゅょゎゕゖ'
    'ァィゥェォッャュョヮヵヶ'
)
# 行末禁則: these may not END a line.
NO_LINE_END = '（「『【〈《〔｛［“‘([{<'
# Breaking right after one of these is the most natural place to break.
GOOD_BREAK_AFTER = '。、，．！？：；」』）'


def _adjust_break(text, start, cut, window):
    """Move a greedy break point to a legal, and if possible natural, position."""
    n = len(text)
    # 1. prefer a natural break within the lookback window
    lo = max(start + 1, cut - window)
    for i in range(cut, lo - 1, -1):
        if i > start and text[i - 1] in GOOD_BREAK_AFTER and \
                (i >= n or text[i] not in NO_LINE_START):
            return i
    # 2. do not orphan an opening bracket / do not start a line with a closer
    guard = 0
    while cut > start + 1 and guard < 8:
        if cut < n and text[cut] in NO_LINE_START:
            cut -= 1
        elif text[cut - 1] in NO_LINE_END:
            cut -= 1
        else:
            break
        guard += 1
    # 3. keep a short quoted span whole
    op = text.rfind('「', start, cut)
    if op > start:
        cl = text.find('」', op)
        if cl != -1 and cl >= cut and (cl - op) <= 24:
            return op
    return cut


def wrap_ja(text, width_in, pt, font=None, window=10):
    """Break Japanese text into lines that fit width_in, honouring 禁則処理.

    A safety net, not the primary mechanism: a slide's key message should carry
    the line breaks a human chose in the storyboard. Use this for body text, and
    to check whether an authored line will hold at a given width.
    """
    mono = (font or '') in MONO_FONTS
    unit = pt / 72.0
    out = []
    for hard in text.split('\n'):
        if not hard:
            out.append('')
            continue
        start, n = 0, len(hard)
        while start < n:
            w, end = 0.0, start
            while end < n:
                nw = w + _em(hard[end], mono) * unit
                if end > start and nw > width_in:
                    break
                w, end = nw, end + 1
            if end >= n:
                out.append(hard[start:])
                break
            cut = _adjust_break(hard, start, end, window)
            out.append(hard[start:cut])
            start = cut
            while start < n and hard[start] == ' ':
                start += 1
    return out


def relative_luminance(hexcolor):
    c = hexcolor.lstrip('#')
    vals = []
    for i in (0, 2, 4):
        v = int(c[i:i + 2], 16) / 255.0
        vals.append(v / 12.92 if v <= 0.04045 else ((v + 0.055) / 1.055) ** 2.4)
    return 0.2126 * vals[0] + 0.7152 * vals[1] + 0.0722 * vals[2]


def contrast_ratio(fg, bg):
    """WCAG contrast ratio. Body text wants >= 4.5, large/bold text >= 3.0."""
    a, b = relative_luminance(fg), relative_luminance(bg)
    lo, hi = min(a, b), max(a, b)
    return (hi + 0.05) / (lo + 0.05)
