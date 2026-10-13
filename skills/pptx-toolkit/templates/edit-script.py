# -*- coding: utf-8 -*-
"""<bNNN wave N / rNNN batch A>: slides 16, 20, 22 -- what this batch does, in one line.

Copy this file next to the project, fill in TARGET and EDITS, run it with python.

It carries the four guards a build script needs, so a run the API cut off half-way is
safe to resume instead of being re-done from memory:

  1. refuses to touch any slide outside TARGET
  2. asserts the slide count and the titles of untouched slides before and after
  3. saves to a temp file and replaces the deck atomically, after EVERY edit
  4. proves every slide outside TARGET is XML-identical afterwards

and it appends a line to PROGRESS after each stage ("saved 20"), which the orchestrator
reads to tell you where to continue. Markdown state files are NOT edited here -- use
techdeck-workflow/scripts/statekit.py for storyboard / review_state / source_index.
"""
import io
import os
import sys
import time

sys.path.insert(0, r'<pptx-toolkit>/scripts')          # ← the real absolute path
from pptx import Presentation                            # noqa: E402
from deckkit import (R, set_runs, set_para, shape_by_id, el_by_id, set_xfrm,   # noqa: E402,F401
                     rrect, textbox, shape, arrow, connector, move_before, move_after,
                     save_backup, load_profile, use_defaults, slide_titles,
                     slide_fingerprints, assert_untouched, save_replace)

PROJECT = r'<project>'                                   # ← the deck project root
DECK = os.path.join(PROJECT, 'slides', 'deck.pptx')
PROGRESS = os.environ.get('PROGRESS_FILE') or os.path.join(PROJECT, 'output', 'build-progress.txt')
TARGET = {16, 20, 22}                                    # 1-based positions this script may touch

PROF = load_profile(os.path.join(PROJECT, 'slides', 'deck-profile.json'))
PAL = PROF['palette']
CJK = PROF['fonts']['cjk']
use_defaults(text=PAL['text'], rule=PAL['rule'])         # ← first, always (light decks)


def log(stage, key=''):
    os.makedirs(os.path.dirname(PROGRESS), exist_ok=True)
    with io.open(PROGRESS, 'a', encoding='utf-8') as f:
        f.write('%s %s %s\n' % (stage, key, time.strftime('%Y-%m-%dT%H:%M:%S')))


def S(prs, n):
    """Slide n (1-based) -- guard 1."""
    if n not in TARGET:
        raise SystemExit('refusing to touch slide %d: not in TARGET %s' % (n, sorted(TARGET)))
    return prs.slides[n - 1]


# ------------------------------------------------------------------ the edits, one function each
# The list EDITS below is also the recovery list: "what should have been written".

def edit_16(prs):
    s = S(prs, 16)
    set_runs(shape_by_id(s, 44), [[R('…', 12, PAL['text'], CJK)]])
    # budget-exempt text gets a role name so check_deck does not count it (visual.uncounted)
    textbox(s, 0.4, 5.2, 6.0, 0.25, [[R('SRC-012 厚生労働省 2026-10-01', 7, PAL['dim'], CJK)]],
            name='src:footer')


def edit_20(prs):
    s = S(prs, 20)
    # ...


EDITS = [('16', edit_16), ('20', edit_20)]


# ------------------------------------------------------------------ run, with the guards
def main():
    save_backup(DECK)
    before = slide_fingerprints(DECK)
    prs = Presentation(DECK)
    n_before = len(prs.slides)
    titles_before = {pos: t for pos, t in slide_titles(prs) if pos not in TARGET}

    done = set()
    if os.path.exists(PROGRESS):
        with io.open(PROGRESS, encoding='utf-8') as f:
            done = {ln.split()[1] for ln in f if ln.startswith('saved ')}

    for key, fn in EDITS:
        if key in done:
            print('skip %s (already saved)' % key)
            continue
        fn(prs)
        log('drawn', key)
        assert len(prs.slides) == n_before, 'slide count changed'           # guard 2
        save_replace(prs, DECK)                                            # guard 3
        assert_untouched(before, slide_fingerprints(DECK), TARGET)         # guard 4
        log('saved', key)

    prs = Presentation(DECK)
    titles_after = {pos: t for pos, t in slide_titles(prs) if pos not in TARGET}
    assert titles_after == titles_before, 'a title outside TARGET changed'   # guard 2
    log('checked')
    print('done: %s — %d untouched slide(s) verified identical'
          % ([k for k, _ in EDITS], n_before - len(TARGET)))


if __name__ == '__main__':
    main()
