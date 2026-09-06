# -*- coding: utf-8 -*-
"""CMT-r001-001 / 003-008 / 012 : text-level fixes on slides 19, 22, 23, 24, 28."""
import copy
import sys
sys.path.insert(0, r'<pptx-toolkit>/scripts')
from pptx import Presentation
from deckkit import (R, set_runs, shape_by_id, el_by_id, set_xfrm, spTree,
                     next_id, _q, frag)

DECK = r'<project>\slides\deck.pptx'
prs = Presentation(DECK)
S = lambda n: prs.slides[n - 1]

MONO = 'Courier New'
SANS = 'Calibri'

# ---------------------------------------------------------------- slide 19
# CMT-r001-001 : 「親は待たず継続」→「親は子を待たず継続」
s19 = S(19)
code = shape_by_id(s19, 7)
for para in code.text_frame.paragraphs:
    for r in para.runs:
        if r.text == '親は待たず継続':
            r.text = '親は子を待たず継続'
out = shape_by_id(s19, 10)
for para in out.text_frame.paragraphs:
    for r in para.runs:
        if r.text.startswith('③ 親は待たず継続'):
            r.text = r.text.replace('親は待たず継続', '親は子を待たず継続')

# ---------------------------------------------------------------- slide 22
s22 = S(22)
# CMT-r001-005 : IO の実行ポリシー「多数まで拡張」を書き直す
set_runs(shape_by_id(s22, 23), [[
    R('blocking ', 8.5, 'ECEAF7', SANS),
    R('用に thread 数を増やせる共有 pool', 8.5, 'ECEAF7', SANS),
]])
# IO カード下部：具体的な上限値をここに逃がす
set_runs(shape_by_id(s22, 24), [[
    R('Default と同じ scheduler を共有（上限は既定 64）→ Slide 27', 7.6, '9C95BC', SANS, i=True),
]])
# CMT-r001-003 : Unconfined の「特殊用途」を具体化
set_runs(shape_by_id(s22, 41), [[
    R('dispatch ', 8.5, 'ECEAF7', SANS),
    R('を待たずその場で実行したい処理（CPU も UI も触らないもの）', 8.5, 'ECEAF7', SANS),
]])
set_runs(shape_by_id(s22, 44), [[
    R('一般のアプリコードでは使わない。テストには ', 7.6, 'FF7597', SANS, b=True),
    R('UnconfinedTestDispatcher', 7.6, 'FF7597', MONO, b=True),
    R(' を使う。', 7.6, 'FF7597', SANS, b=True),
]])
# CMT-r001-004 : 要点バーの「仕組みを理解するため」という価値づけをやめる
set_runs(shape_by_id(s22, 47), [[
    R('通常の選択肢は ', 10.0, '9C95BC', SANS),
    R('Default / IO / Main', 10.0, 'ECEAF7', MONO, b=True),
    R('。', 10.0, '9C95BC', SANS),
    R('Unconfined', 10.0, 'FF7597', MONO, b=True),
    R(' は dispatch を挟まない特殊な Dispatcher ─ 一般のアプリコードでは使わない。',
      10.0, '9C95BC', SANS),
]])

# ---------------------------------------------------------------- slide 23
s23 = S(23)
# CMT-r001-006 : コード注釈 「// ③ → Default」→「// ③ 補完 → Default」
for para in shape_by_id(s23, 7).text_frame.paragraphs:
    for r in para.runs:
        if r.text == '// ③ → Default':
            r.text = '// ③ 補完 → Default'

# CMT-r001-007 : PARENT SCOPE ボックスを子カードと同じ4段構成に揃える
set_xfrm(el_by_id(s23, 28), y=1.62, h=1.00)          # 枠
set_xfrm(el_by_id(s23, 29), y=1.68, h=0.19)          # 見出し PARENT SCOPE
# バッジ「③ 補完」── 子カードの「② 継承」チップを複製して色と文言を変える
chip_src = el_by_id(s23, 36)
chip = copy.deepcopy(chip_src)
chip.find('.//' + _q('p:cNvPr')).set('id', str(next_id(s23)))
chip.find('.//' + _q('p:cNvPr')).set('name', 'Rounded Rectangle Parent Badge')
chip.find('.//' + _q('a:srgbClr')).set('val', 'B89CFF')
spTree(s23).append(chip)
set_xfrm(chip, x=7.28, y=1.94, w=0.66, h=0.22)
for t in chip.iter(_q('a:t')):
    t.text = '③ 補完'
# 「実行 = Default」── 子カードと同じ書式
set_xfrm(el_by_id(s23, 31), x=6.30, y=2.20, w=2.62, h=0.24)
set_runs(shape_by_id(s23, 31), [[
    R('実行 = ', 9.0, '9C95BC', MONO),
    R('Default', 9.0, 'B89CFF', MONO, b=True),
]])
# 補足行（子カードの「親の context をそのまま使用」と同じ位置づけ）
set_xfrm(el_by_id(s23, 30), x=6.30, y=2.44, w=2.62, h=0.18)
# 接続線の起点を新しい枠の下端へ
set_xfrm(el_by_id(s23, 32), y=2.66, h=0.62)
set_xfrm(el_by_id(s23, 33), y=2.66, h=0.62)

# ---------------------------------------------------------------- slide 24
# CMT-r001-008 : 「補足」を「何の実例か → 何が起きるか → だから何が言えるか」に書き直す
s24 = S(24)
set_xfrm(el_by_id(s24, 50), x=1.46, y=4.37, w=7.90, h=0.20)
set_runs(shape_by_id(s24, 50), [[
    R('val ViewModel.', 7.8, '9C95BC', MONO),
    R('viewModelScope', 7.8, '34E2B0', MONO, b=True),
    R(': CoroutineScope get() = ... CoroutineScope(SupervisorJob() + ', 7.8, '9C95BC', MONO),
    R('Dispatchers.Main.immediate', 7.8, '34E2B0', MONO, b=True),
    R(')', 7.8, '9C95BC', MONO),
]])
set_xfrm(el_by_id(s24, 51), x=1.46, y=4.59, w=7.90, h=0.32)
set_runs(shape_by_id(s24, 51), [[
    R('② 継承 の実務例 ─ scope 自身が ', 7.5, 'ECEAF7', SANS),
    R('Dispatchers.Main.immediate', 7.5, '34E2B0', MONO),
    R(' を持つので、', 7.5, 'ECEAF7', SANS),
    R('viewModelScope.launch { }', 7.5, 'ECEAF7', MONO),
    R(' は明示しなくても Main で動く。', 7.5, 'ECEAF7', SANS),
    R('継承元は launch のレシーバで決まり、コード上の入れ子とは無関係（', 7.5, '6E6A8A', SANS, i=True),
    R('lifecycleScope', 7.5, '34E2B0', MONO, i=True),
    R(' も同様）。', 7.5, '6E6A8A', SANS, i=True),
]])

# ---------------------------------------------------------------- slide 28
# CMT-r001-012 : 要点①の主語を suspend から Dispatcher へ
s28 = S(28)
set_runs(shape_by_id(s28, 7), [[
    R('実行 thread を決めるのは Dispatcher', 12.5, 'ECEAF7', SANS, b=True),
]])
set_runs(shape_by_id(s28, 8), [[
    R('suspend', 9.0, 'B89CFF', MONO, b=True),
    R(' は中断・再開できる形を作るだけ ─ どの thread で動くかは context の Dispatcher が決める。',
      9.0, '9C95BC', SANS),
]])
# 要点②：①（どの thread か）との役割分担をはっきりさせる
set_runs(shape_by_id(s28, 11), [[
    R('Dispatcher が scheduling するのは 初回実行 と 再開', 12.5, 'ECEAF7', SANS, b=True),
]])

prs.save(DECK)
print('saved')
