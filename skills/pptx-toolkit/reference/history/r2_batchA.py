# -*- coding: utf-8 -*-
"""r002 Batch A: reaction slides 16, 20, 22, 23, 24, 25 (text/label)."""
import sys
sys.path.insert(0, r'<pptx-toolkit>/scripts')
from pptx import Presentation
from deckkit import R, set_runs, shape_by_id, rrect, textbox, move_before, el_by_id

DECK = r'<project>\slides\deck.pptx'
prs = Presentation(DECK)
S = lambda n: prs.slides[n - 1]
MONO, SANS, GOTH = 'Courier New', 'Calibri', 'Yu Gothic'
LIGHT, MUTED, DIM = 'ECEAF7', '9C95BC', '6E6A8A'
GREEN, AMBER, PURPLE, PINK = '34E2B0', 'FFB454', 'B89CFF', 'FF7597'

# ---------------------------------------------------------------- slide 16
# CMT-r002-007: bullet 3 references cut Unconfined/limitedParallelism
set_runs(shape_by_id(S(16), 10), [[
    R('Default / IO / Main の使い分けと、継承・', 12.0, 'F4F0FF', GOTH),
    R('withContext', 12.0, 'F4F0FF', GOTH),
    R(' の往復を押さえる', 12.0, 'F4F0FF', GOTH),
]])

# ---------------------------------------------------------------- slide 20
s20 = S(20)
# CMT-r002-002: drop the "注: " prefix
set_runs(shape_by_id(s20, 104), [[
    R('CoroutineStart.UNDISPATCHED', 7.5, PINK, MONO, b=True),
    R(' は intercept / dispatch を経ず、呼び出し元 thread で最初の suspension まで直接 '
      'invokeSuspend する（初回 dispatch の例外）。', 7.5, MUTED, SANS),
]])
# CMT-r002-001a: ⑦ label — say plainly that the app code runs here
set_runs(shape_by_id(s20, 97), [[
    R('⑦ ', 8.5, MUTED, SANS, b=True),
    R('invokeSuspend()', 8.5, LIGHT, MONO, b=True),
    R('＝あなたの ', 7.5, MUTED, SANS),
    R('launch{ }', 7.5, LIGHT, MONO),
    R(' の続き', 7.5, MUTED, SANS),
]])
# widen that label box so the longer text fits
el_by_id(s20, 96)  # backing autoshape (unused width match); textbox is 97
from deckkit import set_xfrm
set_xfrm(el_by_id(s20, 97), x=6.05, w=3.45)
set_xfrm(el_by_id(s20, 96), x=6.05, w=3.45)
# CMT-r002-001b: answer "how does it resume / where does my code run" in the empty left column
box = rrect(s20, 0.55, 3.30, 2.05, 1.30, fill='1E1838', line='3C3463', lw=9525, adj=6000)
tb = textbox(s20, 0.68, 3.42, 1.83, 1.06, [
    [R('この図全体が「再開1回」の経路', 8.0, AMBER, SANS, b=True)],
    [R('あなたの ', 7.6, MUTED, SANS), R('launch{ }', 7.6, LIGHT, MONO),
     R(' の中身は', 7.6, MUTED, SANS)],
    [R('⑥⑦ の時点で worker thread 上を走る。', 7.6, MUTED, SANS)],
    [R('再び suspend すれば、次の ', 7.6, MUTED, SANS)],
    [R('resume で ① から同じ経路を繰り返す。', 7.6, MUTED, SANS)],
], algn='l', spc_pct=108)
move_before(s20, box, el_by_id(s20, 49))
move_before(s20, tb, el_by_id(s20, 49))

# ---------------------------------------------------------------- slide 22
# CMT-r002-003: cut the UnconfinedTestDispatcher sentence
set_runs(shape_by_id(S(22), 44), [[
    R('一般のアプリコードでは使わない。', 7.6, PINK, SANS, b=True),
]])

# ---------------------------------------------------------------- slide 23
# CMT-r002-004: flowchart result box ①「明示」→「上書き」
set_runs(shape_by_id(S(23), 14), [[R('上書き', 9.3, AMBER, SANS, b=True)]])

# ---------------------------------------------------------------- slide 24
s24 = S(24)
# CMT-r002-006 (consistency): rule label ①「明示あり」→「上書き」
for sh in s24.shapes:
    if sh.has_text_frame and sh.text_frame.text.replace('\n', '') == '①明示あり':
        set_runs(sh, [[R('①', 12.5, AMBER, SANS, b=True)], [R('上書き', 9.0, AMBER, SANS, b=True)]])
        print('slide24 rule label id', sh.shape_id, 'updated')
# CMT-r002-005: viewModelScope — explicit at scope construction, inherited at launch
set_runs(shape_by_id(s24, 51), [[
    R('scope 構築時に ', 7.5, LIGHT, SANS),
    R('①明示', 7.5, AMBER, SANS, b=True),
    R(' された Dispatcher を、launch が ', 7.5, LIGHT, SANS),
    R('②継承', 7.5, GREEN, SANS, b=True),
    R(' する例 ─ ', 7.5, LIGHT, SANS),
    R('Dispatchers.Main.immediate', 7.5, '34E2B0', MONO),
    R(' は viewModelScope を作った時点で選ばれる。だから ', 7.5, LIGHT, SANS),
    R('viewModelScope.launch { }', 7.5, LIGHT, MONO),
    R(' は指定なしでも Main（既定 Default ではない）。', 7.5, LIGHT, SANS),
    R('明示すれば上書き可（', 7.5, DIM, SANS, i=True),
    R('lifecycleScope', 7.5, '34E2B0', MONO, i=True),
    R(' も同様）。', 7.5, DIM, SANS, i=True),
]])

# ---------------------------------------------------------------- slide 25
# CMT-r002-006: A/C card heading "caller context (Default)" → "Default context"
for sid in (13, 23):
    sh = shape_by_id(S(25), sid)
    for para in sh.text_frame.paragraphs:
        for r in para.runs:
            if r.text == 'caller context (Default)':
                r.text = 'Default context'

prs.save(DECK)
print('Batch A saved')
