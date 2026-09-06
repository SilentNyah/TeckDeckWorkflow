# -*- coding: utf-8 -*-
"""r002 Batch C: slides 33, 34, 35, 38, 40."""
import sys
sys.path.insert(0, r'<pptx-toolkit>/scripts')
from pptx import Presentation
from deckkit import (R, set_runs, set_para, shape_by_id, el_by_id, set_xfrm,
                     connector, textbox, move_before, spTree, _q)

DECK = r'<project>\slides\deck.pptx'
prs = Presentation(DECK)
S = lambda n: prs.slides[n - 1]
MONO, SANS = 'Courier New', 'Calibri'
LIGHT, MUTED, DIM = 'ECEAF7', '9C95BC', '6E6A8A'
GREEN, AMBER, PURPLE, PINK = '34E2B0', 'FFB454', 'B89CFF', 'FF7597'


def set_line_color(el, color):
    ln = el.find('.//' + _q('a:ln'))
    sc = ln.find('.//' + _q('a:srgbClr'))
    sc.set('val', color)


def replace_run_text(shape, old, new):
    for pa in shape.text_frame.paragraphs:
        for r in pa.runs:
            if r.text == old:
                r.text = new
                return True
    return False


# ================================================================ slide 33
s33 = S(33)
# CMT-r002-012: lead — the parent doesn't need to join the CHILD (main's join is separate)
set_runs(shape_by_id(s33, 4), [[
    R('子を追跡する必要も、', 11.0, MUTED, SANS),
    R('親が子を明示 join する必要もない', 11.0, LIGHT, SANS, b=True),
    R(' ─ 親は全子の完了を暗黙に待つ。', 11.0, MUTED, SANS),
]])
# code comment on request.join() — say plainly this is main's wait, not the parent-child mechanism
replace_run_text(shape_by_id(s33, 7), '// request と全子の完了を待つ',
                 '// これは main の待ち（親子の join とは別）')
# CMT-r002-013: name the figure axes
set_runs(shape_by_id(s33, 11), [[R('Job の視点 ─ 横軸＝時間、縦＝親と各子', 8.5, MUTED, SANS, b=True)]])
set_para(shape_by_id(s33, 26), 0, [R('バーの幅＝実行時間。body 完了後も最後の子（child 2）まで', 8.19, MUTED, SANS)])
# a small time arrow under the child bars
connector(s33, 6.30, 3.45, 1.75, 0, color=DIM, lw=9525, tail='triangle')
tb = textbox(s33, 6.30, 3.46, 1.75, 0.15, [[R('時間 →', 7.0, DIM, SANS, i=True)]], algn='r')

# ================================================================ slide 34
s34 = S(34)
# CMT-r002-014 sp6: parent box shows it is a Job
set_runs(shape_by_id(s34, 6), [
    [R('request', 9.5, LIGHT, MONO, b=True)],
    [R('＝ 親 Job', 7.0, MUTED, SANS)],
])
# sp10: swap so job1 (survives) is LEFT, job2 (cancelled) is RIGHT
set_xfrm(el_by_id(s34, 15), x=0.80)   # job1 ○  → left
set_xfrm(el_by_id(s34, 12), x=2.95)   # job2 ✕  → right
set_xfrm(el_by_id(s34, 16), x=0.55)   # "リンク無し → 届かない" (job1) → left
set_xfrm(el_by_id(s34, 13), x=2.70)   # "通常の子 → cancel が届く" (job2) → right
# redraw arrows: dashed→job1(left, severed), solid→job2(right, cancel reaches)
for sid in (8, 9):
    el = el_by_id(s34, sid)
    el.getparent().remove(el)
connector(s34, 2.35, 2.06, 1.275, 0.56, color=PURPLE, lw=19050, tail='triangle')          # →job2 right
connector(s34, 1.475, 2.06, 0.875, 0.56, color=DIM, lw=12700, dash='dash',
          tail='triangle', flipH=True)                                                     # →job1 left
set_xfrm(el_by_id(s34, 10), x=0.52, y=2.20)   # "launch(Job()) ＝切断" (job1) → left
# sp16: make code comment colors match the diagram (job1 survives=green, job2 cancelled=pink)
from pptx.dml.color import RGBColor
for pa in shape_by_id(s34, 22).text_frame.paragraphs:
    for r in pa.runs:
        if r.text == '// リンク切断':
            r.font.color.rgb = RGBColor.from_string(GREEN)
        elif r.text == '// 通常の子':
            r.font.color.rgb = RGBColor.from_string(PINK)
# sp18: rewrite the "cancel(cause)..." note
set_xfrm(el_by_id(s34, 17), h=0.74)
set_runs(shape_by_id(s34, 18), [
    [R('↑ 親へは伝播しない ─ 子を cancel しても、それは', 8.19, LIGHT, SANS, b=True)],
    [R('CancellationException', 8.19, PINK, MONO), R('（＝通常の停止）。失敗ではないので親へは上がらない。', 8.19, MUTED, SANS)],
])

# ================================================================ slide 35
s35 = S(35)
# CMT-r002-015 sp22: explain the NonCancellable line
set_runs(shape_by_id(s35, 22), [[
    R('finally で cleanup', 8.6, PURPLE, SANS, b=True),
    R('  cancel 後の suspend は即 throw ─ 片付けで suspend するなら ', 6.9, MUTED, SANS),
    R('withContext(NonCancellable)', 6.9, MUTED, MONO),
]])
# CMT-r002-015 sp11 + slide-level: 3-API heading + what CancellationException is
set_xfrm(el_by_id(s35, 11), w=5.02, h=0.48)
set_runs(shape_by_id(s35, 11), [
    [R('CancellationException', 8.0, PINK, MONO, b=True),
     R(' ＝ coroutine を巻き戻す通常の合図（失敗ではない）', 8.0, LIGHT, SANS)],
    [R('協調 cancel の道具 ─ ', 8.0, MUTED, SANS),
     R('isActive', 8.0, GREEN, MONO), R('（確認）／ ', 8.0, MUTED, SANS),
     R('ensureActive()', 8.0, GREEN, MONO), R('（throw）／ ', 8.0, MUTED, SANS),
     R('yield()', 8.0, GREEN, MONO), R('（check＋譲る）', 8.0, MUTED, SANS)],
])

# ================================================================ slide 38
s38 = S(38)
# CMT-r002-018 sp15/sp26: make child A (left) the failing origin in both cards
set_runs(shape_by_id(s38, 10), [[R('child A 失敗', 8.5, AMBER, MONO, b=True)]])   # coroutineScope A
set_runs(shape_by_id(s38, 13), [[R('child B ✕', 8.5, PINK, MONO, b=True)]])       # coroutineScope B
set_runs(shape_by_id(s38, 21), [[R('child A 失敗', 8.5, AMBER, MONO, b=True)]])   # supervisorScope A
set_runs(shape_by_id(s38, 24), [[R('child B ○', 8.5, GREEN, MONO, b=True)]])      # supervisorScope B
set_line_color(el_by_id(s38, 9), AMBER)
set_line_color(el_by_id(s38, 12), PINK)
set_line_color(el_by_id(s38, 20), AMBER)
set_line_color(el_by_id(s38, 23), GREEN)
set_line_color(el_by_id(s38, 11), AMBER)   # arrow →A (coroutineScope)
set_line_color(el_by_id(s38, 14), PINK)    # arrow →B
set_line_color(el_by_id(s38, 22), AMBER)   # arrow →A (supervisorScope)
set_line_color(el_by_id(s38, 25), GREEN)   # arrow →B
set_para(shape_by_id(s38, 15), 0, [R('A の失敗 → scope が cancel → B も ✕', 9.0, PINK, SANS, b=True)])
set_para(shape_by_id(s38, 26), 0, [R('A の失敗は A だけ ─ scope と B は継続', 9.0, GREEN, SANS, b=True)])
# CMT-r002-018 sp4: clarify it only stops UPWARD propagation
set_runs(shape_by_id(s38, 4), [[
    R('SupervisorJob', 11.0, GREEN, MONO, b=True),
    R(' / ', 11.0, GREEN, MONO, b=True),
    R('supervisorScope', 11.0, GREEN, MONO, b=True),
    R(' が止めるのは失敗の上方向（子→親）伝播だけ ─ 親→子の下方向 cancel は通常どおり。', 11.0, MUTED, SANS),
]])

# ================================================================ slide 40
s40 = S(40)
# CMT-r002-020 sp4: replace jargon 字句的
set_runs(shape_by_id(s40, 4), [[
    R('その場かぎり（suspend 関数の中だけ）', 11.0, MUTED, SANS),
    R('／ 長命（lifecycle に紐づけ）／ bridge の3系統 ─ ', 11.0, MUTED, SANS),
    R('GlobalScope は原則使わない', 11.0, LIGHT, SANS, b=True),
    R('。', 11.0, MUTED, SANS),
]])
set_runs(shape_by_id(s40, 11), [[R('その場（suspend 内）', 8.0, LIGHT, SANS)]])
set_runs(shape_by_id(s40, 16), [[R('その場（suspend 内）', 8.0, LIGHT, SANS)]])
# CMT-r002-020 sp38: name the axes of the lifecycle figure
set_runs(shape_by_id(s40, 38), [[
    R('entity の寿命（横軸＝時間）と scope 配下の coroutine（縦）が一緒に終わる', 8.0, MUTED, SANS)]])

prs.save(DECK)
print('Batch C saved')
