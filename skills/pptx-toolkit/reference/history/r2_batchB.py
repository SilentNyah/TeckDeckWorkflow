# -*- coding: utf-8 -*-
"""r002 Batch B: slide 24 (①上書き consistency), 31 (isActive + Job arrow),
32 (box both internal states + 区別 rewrite), 36 (recursion)."""
import copy
import sys
sys.path.insert(0, r'<pptx-toolkit>/scripts')
from pptx import Presentation
from deckkit import (R, set_runs, shape_by_id, el_by_id, set_xfrm, spTree,
                     next_id, connector, move_after, _q, frag, _run_xml)

DECK = r'<project>\slides\deck.pptx'
prs = Presentation(DECK)
S = lambda n: prs.slides[n - 1]
MONO, SANS = 'Courier New', 'Calibri'
LIGHT, MUTED, DIM = 'ECEAF7', '9C95BC', '6E6A8A'
GREEN, AMBER, PURPLE, PINK = '34E2B0', 'FFB454', 'B89CFF', 'FF7597'


def set_para(shape, pidx, runs):
    """Replace runs of a single paragraph, keeping the rest of the text frame."""
    p_el = shape.text_frame.paragraphs[pidx]._p
    for r in p_el.findall(_q('a:r')):
        p_el.remove(r)
    for rr in runs:
        p_el.append(frag(_run_xml(rr))[0])


# ================================================================ slide 24
s24 = S(24)
# subtitle ①明示 → ①上書き
for para in shape_by_id(s24, 4).text_frame.paragraphs:
    for r in para.runs:
        if r.text == '①明示 ／ ②継承 ／ ③補完':
            r.text = '①上書き ／ ②継承 ／ ③補完'
# 要点 ①明示 → ①上書き
for para in shape_by_id(s24, 54).text_frame.paragraphs:
    for r in para.runs:
        if r.text == '①明示 → ②継承 → ③補完':
            r.text = '①上書き → ②継承 → ③補完'
# 補足 rewrite (CMT-r002-005): construction = explicit; launch = ②継承
set_runs(shape_by_id(s24, 51), [[
    R('viewModelScope は構築時に ', 7.5, LIGHT, SANS),
    R('Dispatchers.Main.immediate', 7.5, GREEN, MONO),
    R(' を明示的に持つ ─ だから ', 7.5, LIGHT, SANS),
    R('viewModelScope.launch { }', 7.5, LIGHT, MONO),
    R(' は指定なしでも ', 7.5, LIGHT, SANS),
    R('②継承', 7.5, GREEN, SANS, b=True),
    R(' で Main（', 7.5, LIGHT, SANS),
    R('③補完', 7.5, PURPLE, SANS, b=True),
    R(' の既定 Default ではない）。launch で明示すれば ', 7.5, LIGHT, SANS),
    R('①上書き', 7.5, AMBER, SANS, b=True),
    R(' 可（', 7.5, DIM, SANS, i=True),
    R('lifecycleScope', 7.5, GREEN, MONO, i=True),
    R(' も同様）。', 7.5, DIM, SANS, i=True),
]])

# ================================================================ slide 31
s31 = S(31)
# CMT-r002-008: fix circular "isActive ＝ その isActive の短縮" (paragraph P9)
set_para(shape_by_id(s31, 6), 9, [
    R('isActive', 9.5, GREEN, MONO),
    R(' ＝ ', 9.0, MUTED, SANS),
    R('coroutineContext[Job]?.isActive', 7.6, GREEN, MONO),
    R(' の短縮', 9.0, MUTED, SANS),
])
# also make P8 a touch clearer: 自分の Job → いま自分がいる Job
set_para(shape_by_id(s31, 6), 8, [
    R('coroutineContext[Job]', 9.5, GREEN, MONO),
    R(' ─ いま自分がいる Job', 9.0, MUTED, SANS),
])
# CMT-r002-009: redraw the parent-child link so it connects Job box → Job box
for sid in (19, 20, 21):
    el = el_by_id(s31, sid)
    el.getparent().remove(el)
# Job=C box (id17) bottom-center (8.52,2.60) → down, left, up into Job=P (id10) bottom-center (5.80,2.60)
connector(s31, 8.52, 2.60, 0, 0.42, color=PURPLE, lw=12700)                 # down from C
connector(s31, 5.80, 3.02, 2.72, 0, color=PURPLE, lw=12700)                 # left
connector(s31, 5.80, 2.60, 0, 0.42, color=PURPLE, lw=12700, head='triangle')  # up into P
set_xfrm(el_by_id(s31, 22), y=3.06)   # label just under the horizontal segment

# ================================================================ slide 32
s32 = S(32)
# CMT-r002-010: box the Cancelling row too (both internal states), matching id42
box = copy.deepcopy(el_by_id(s32, 42))
box.find('.//' + _q('p:cNvPr')).set('id', str(next_id(s32)))
box.find('.//' + _q('p:cNvPr')).set('name', 'Rounded Rectangle Cancelling HL')
spTree(s32).append(box)
set_xfrm(box, y=2.72)               # Cancelling row (id47 at y=2.74)
move_after(s32, box, el_by_id(s32, 42))
# CMT-r002-011: clarify "区別は原因例外" in the 要点 bar
set_runs(shape_by_id(s32, 62), [[
    R('body 完了 ≠ Job 完了。', 9.5, LIGHT, SANS, b=True),
    R('　normal cancellation と失敗はどちらも Cancelling へ入る ─ 違いは中断の原因が ',
      9.5, MUTED, SANS),
    R('CancellationException', 9.5, LIGHT, MONO),
    R(' か、それ以外の例外かだけ（Slide 34 / 36）。', 9.5, MUTED, SANS),
]])

# ================================================================ slide 36
s36 = S(36)
# CMT-r002-016: make the upward recursion explicit
set_runs(shape_by_id(s36, 4), [[
    R('CancellationException 以外の失敗は ', 11.0, MUTED, SANS),
    R('root まで再帰的に上がって親を cancel し', 11.0, LIGHT, SANS, b=True),
    R('、その配下の兄弟も全て cancel される（既定・上書き不可）。', 11.0, MUTED, SANS),
]])
set_runs(shape_by_id(s36, 17), [[R('方向: root まで ↑、その後 ↓', 8.5, AMBER, SANS, b=True)]])

prs.save(DECK)
print('Batch B saved')
