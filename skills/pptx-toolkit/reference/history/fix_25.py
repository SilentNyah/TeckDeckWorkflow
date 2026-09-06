# -*- coding: utf-8 -*-
"""CMT-r001-009 / 010 / 011 : slide 25, plus a tweak to slide 22's Unconfined note."""
import sys
sys.path.insert(0, r'<pptx-toolkit>/scripts')
from pptx import Presentation
from deckkit import (R, set_runs, shape_by_id, el_by_id, rrect, textbox,
                     move_before, spTree, _q)

DECK = r'<project>\slides\deck.pptx'
prs = Presentation(DECK)
S = lambda n: prs.slides[n - 1]

MONO = 'Courier New'
SANS = 'Calibri'
PURPLE = 'B89CFF'
AMBER = 'FFB454'
PINK = 'FF7597'

# --------------------------------------------------- slide 22 (fit-up only)
set_runs(shape_by_id(S(22), 44), [[
    R('一般のアプリコードでは使わない。テスト用途には ', 7.2, PINK, SANS, b=True),
    R('UnconfinedTestDispatcher', 7.2, PINK, MONO, b=True),
    R(' が別にある。', 7.2, PINK, SANS, b=True),
]])

# ---------------------------------------------------------------- slide 25
s25 = S(25)

# CMT-r001-010 : trace ラベルを Dispatcher 名で統一（A: caller → A: Default 等）
rename = {'"A: caller"': '"A: Default"', '"C: caller"': '"C: Default"',
          'trace("A: caller")': 'trace("A: Default")',
          'trace("C: caller")': 'trace("C: Default")',
          'A: caller  ': 'A: Default ', 'C: caller  ': 'C: Default '}
for sid in (7, 10, 13, 23):
    for para in shape_by_id(s25, sid).text_frame.paragraphs:
        for r in para.runs:
            if r.text in rename:
                r.text = rename[r.text]

# CMT-r001-009 : 実行 context ごとの帯をコードの背面に敷く
#   コード1行目 "runBlocking {" の中心 y=1.98in、行送り 0.158in
BANDS = [
    # (先頭行, 末尾行, 色, ラベル)
    (3, 4, PURPLE, 'Default'),
    (5, 6, AMBER,  'IO'),
    (7, 8, PURPLE, 'Default'),
]
code_el = el_by_id(s25, 7)
for first, last, color, label in BANDS:
    top = 1.98 + (first - 1) * 0.158 - 0.079
    bottom = 1.98 + (last - 1) * 0.158 + 0.079
    band = rrect(s25, 0.78, top, 4.16, bottom - top,
                 fill=color, fill_alpha=13, line=color, lw=9525, adj=6000)
    move_before(s25, band, code_el)
    tb = textbox(s25, 3.72, top, 1.18, bottom - top,
                 [[R(label, 7.0, color, MONO, b=True)]], algn='r')
    move_before(s25, tb, code_el)
# 帯の意味を1行で説明（コード枠内・最下行の下）
note = textbox(s25, 0.78, 3.24, 4.16, 0.16,
               [[R('帯 = その行が動く Dispatcher ─ ', 6.8, '6E6A8A', SANS, i=True),
                 R('Default → IO → Default', 6.8, '9C95BC', MONO, b=True),
                 R(' の往復', 6.8, '6E6A8A', SANS, i=True)]], algn='l')
move_before(s25, note, code_el)

# CMT-r001-011 : 「要点」→「注意」（Slide 23 の規約に合わせてピンクへ）
bar = el_by_id(s25, 25)
for c in bar.iter(_q('a:srgbClr')):
    if c.get('val') == AMBER:
        c.set('val', PINK)
chip = el_by_id(s25, 26)
for c in chip.find('.//' + _q('a:solidFill')).iter(_q('a:srgbClr')):
    c.set('val', PINK)
for t in shape_by_id(s25, 26).text_frame.paragraphs[0].runs:
    if t.text == '要点':
        t.text = '注意'

prs.save(DECK)
print('saved')
